"""Genera Documentacion/catalogo_fabric_linaje.xlsx desde un repositorio Fabric Git-integrado."""
from pathlib import Path
import csv, json, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Documentacion' / 'catalogo_fabric_linaje.xlsx'
INVENTORY = ROOT / 'Documentacion' / 'objetos_fisicos.csv'
RED, DARK, LIGHT, BORDER = 'B5121B', '151515', 'F4F5F7', 'D9DDE3'
LAKEHOUSES = {}

def classify(name):
    x = name.lower()
    if x.startswith('br_') or 'bronce' in x: return 'Bronze'
    if x.startswith('sl_') or 'silver' in x: return 'Silver'
    if x.startswith('gd_') or 'gold' in x: return 'Gold'
    return 'Sin clasificar'

def extract():
    artifacts, relations = [], []
    for path in ROOT.rglob('*'):
        if not path.is_dir(): continue
        suffix = path.suffix
        kinds = {'.Dataflow':'Dataflow Gen2','.DataPipeline':'Pipeline','.Notebook':'Notebook','.Lakehouse':'Lakehouse'}
        if suffix not in kinds: continue
        name, kind = path.stem, kinds[suffix]
        folder = str(path.parent.relative_to(ROOT)) if path.parent != ROOT else '(raíz)'
        artifacts.append([name, kind, folder, classify(name), 'Pendiente', ''])
        if kind == 'Lakehouse': LAKEHOUSES[path.name] = name
        if kind == 'Dataflow Gen2':
            m = path / 'mashup.pq'
            if not m.exists(): continue
            text = m.read_text(encoding='utf-8', errors='ignore')
            source = ' + '.join(x for x, token in [('ODBC / Informix','Odbc.Query'),('SQL Server','Sql.Database'),('SharePoint','SharePoint.Files')] if token in text) or 'Lakehouse / otras consultas'
            # Physical target table is found inside each DataDestination shared query.
            destinations = re.findall(r'QueryName\s*=\s*"?([^,\]\"]+)', text)
            for dest in sorted(set(destinations)):
                block = re.search(r'shared\s+' + re.escape(dest) + r'\s*=\s*let(.*?)(?=\nshared\s|\Z)', text, re.S)
                table = re.search(r'Id\s*=\s*"([^"]+)"\s*,\s*ItemKind\s*=\s*"Table"', block.group(1)) if block else None
                target = table.group(1) if table else dest.replace('_DataDestination','')
                relations.append([name, 'produce', target, source, 'Confirmado por definición', 'Pendiente'])
        elif kind == 'Pipeline':
            p = path / 'pipeline-content.json'
            if p.exists():
                text = p.read_text(encoding='utf-8', errors='ignore')
                for item in sorted(set(re.findall(r'"itemName"\s*:\s*"([^"]+)', text))):
                    relations.append([name, 'ejecuta', item, 'Pipeline', 'Confirmado por definición', 'Pendiente'])
    return sorted(artifacts), sorted(relations)

def inventory():
    if not INVENTORY.exists(): return []
    with INVENTORY.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def setup(ws, title, widths):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    c = ws.cell(1,1,title); c.fill = PatternFill('solid', fgColor=RED); c.font = Font(bold=True,color='FFFFFF',size=16); c.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width = w

def add_table(ws, start_row, headers, rows, name):
    for j,h in enumerate(headers,1):
        c=ws.cell(start_row,j,h); c.fill=PatternFill('solid',fgColor=DARK); c.font=Font(bold=True,color='FFFFFF'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for r,row in enumerate(rows,start_row+1):
        for c,val in enumerate(row,1):
            cell=ws.cell(r,c,val); cell.alignment=Alignment(vertical='top',wrap_text=True); cell.border=Border(bottom=Side(style='thin',color=BORDER))
    end = start_row + len(rows)
    if rows:
        tab=Table(displayName=name, ref=f'A{start_row}:{chr(64+len(headers))}{end}')
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tab)
    ws.freeze_panes=f'A{start_row+1}'

def main():
    artifacts, relations = extract(); objects = inventory()
    wb=Workbook(); wb.remove(wb.active)
    summary=wb.create_sheet('Resumen'); setup(summary,'Catálogo de capas y linaje — Fabric',[32,18,4,40,28,28])
    summary['A3']='Generado automáticamente desde el repositorio Git y el inventario físico versionado.'; summary['A3'].font=Font(italic=True,color='666666')
    metrics=[['Indicador','Valor'],['Objetos físicos inventariados',len(objects)],['Objetos / vistas Bronze',sum(x['capa_fisica']=='Bronze' for x in objects)],['Objetos / vistas Silver',sum(x['capa_fisica']=='Silver' for x in objects)],['Objetos / vistas Gold',sum(x['capa_fisica']=='Gold' for x in objects)],['Artefactos versionados',len(artifacts)],['Relaciones extraídas',len(relations)]]
    for r,row in enumerate(metrics,5):
        for c,val in enumerate(row,1): summary.cell(r,c,val)
    for cell in summary[5]: cell.fill=PatternFill('solid',fgColor=DARK); cell.font=Font(bold=True,color='FFFFFF')
    summary['D5']='Regla'; summary['E5']='Estado';
    for c in ('D5','E5'): summary[c].fill=PatternFill('solid',fgColor=DARK); summary[c].font=Font(bold=True,color='FFFFFF')
    rules=[['Prefijos','br_, sl_ y gd_ deben coincidir con la capa física.'],['SQL','Actualizar objetos_fisicos.csv al crear una vista o tabla manual.'],['Git','No editar el xlsx maestro; se reemplaza en cada ejecución.']]
    for r,row in enumerate(rules,6): summary.cell(r,4,row[0]); summary.cell(r,5,row[1])

    ws=wb.create_sheet('Artefactos Git'); setup(ws,'Artefactos versionados en Git',[36,20,42,18,20,45]); add_table(ws,3,['Artefacto','Tipo','Carpeta Git','Capa sugerida','Dueño','Observación'],artifacts,'ArtefactosGit')
    ws=wb.create_sheet('Linaje extraído'); setup(ws,'Linaje reconstruido desde Dataflows y Pipelines',[40,16,44,26,28,20]); add_table(ws,3,['Origen','Relación','Destino','Evidencia / fuente','Confianza','Validación'],relations,'LinajeExtraido')
    objrows=[]
    for x in objects:
        logical=classify(x['objeto']); note='La capa física no coincide con el prefijo lógico' if logical != 'Sin clasificar' and logical != x['capa_fisica'] else ''
        objrows.append([x['capa_fisica'],x['esquema'],x['objeto'],x['tipo_objeto'],logical,'Pendiente','Pendiente',note])
    ws=wb.create_sheet('Objetos físicos'); setup(ws,'Inventario de tablas y vistas por Lakehouse',[16,14,44,16,20,20,18,52]); add_table(ws,3,['Capa física','Esquema','Objeto','Tipo','Capa por prefijo','Dueño','Criticidad','Observación automática'],objrows,'ObjetosFisicos')
    ws=wb.create_sheet('Reglas de gobierno'); setup(ws,'Reglas mínimas de gobierno',[25,25,70,22]); add_table(ws,3,['Elemento','Estándar','Regla','Estado'],[['Capa','Bronze / Silver / Gold','Una sola capa física por etapa; las excepciones se documentan.','Pendiente'],['Vistas SQL','vw_','Guardar el DDL en Git para completar linaje de tablas internas.','Pendiente'],['Frecuencia','Documentada','Cada Dataflow/Pipeline debe tener frecuencia, dueño y SLA.','Pendiente']], 'ReglasGobierno')
    OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
    print(f'Catálogo generado: {OUT}')

if __name__ == '__main__': main()
